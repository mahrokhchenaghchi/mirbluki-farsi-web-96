# -*- coding: utf-8 -*-
"""Build Erfan's exercise package as a fully-RTL Word (.docx) document."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD = 'erfan-kamalzadeh-exercises.md'
OUT = 'Erfan-Kamalzadeh-Exercises.docx'
FONT = 'Vazirmatn'

NAVY   = (31, 58, 95)
BLUE   = (44, 74, 124)
GOLDD  = (166, 124, 46)
GRAY   = (110, 105, 98)
DARK   = (40, 40, 45)
WHITE  = (255, 255, 255)
CARD   = (46, 58, 72)
GOLDTX = (232, 182, 74)

FILL_TITLE  = 'E8EEF7'   # light blue  - exercise headings
FILL_SECT   = 'F3E9D2'   # light gold  - section headings
FILL_NOTE   = 'FDF6E3'   # cream       - intro note
FILL_WARN   = 'FDECEA'   # light red   - clinical warning
FILL_DISC   = 'F2F2F2'   # light gray  - disclaimer
FILL_COVER  = 'FBF7EC'   # cover box

# rPr children that must come AFTER a given element (schema order)
_AFTER_BCS  = ('w:i','w:iCs','w:caps','w:smallCaps','w:strike','w:dstrike','w:outline','w:shadow','w:emboss','w:imprint','w:noProof','w:snapToGrid','w:vanish','w:webHidden','w:color','w:spacing','w:w','w:kern','w:position','w:sz','w:szCs','w:highlight','w:u','w:effect','w:bdr','w:shd','w:fitText','w:vertAlign','w:rtl','w:cs','w:em','w:lang','w:eastAsianLayout','w:specVanish','w:oMath')
_AFTER_ICS  = _AFTER_BCS[2:]
_AFTER_SZCS = ('w:highlight','w:u','w:effect','w:bdr','w:shd','w:fitText','w:vertAlign','w:rtl','w:cs','w:em','w:lang','w:eastAsianLayout','w:specVanish','w:oMath')
_AFTER_RTL  = ('w:cs','w:em','w:lang','w:eastAsianLayout','w:specVanish','w:oMath')
_AFTER_BIDI = ('w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing','w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection','w:textAlignment','w:outlineLvl','w:divId','w:cnfStyle','w:rPr','w:sectPr','w:pPrChange')
_AFTER_SHD  = ('w:tabs','w:suppressAutoHyphens','w:kinsoku','w:wordWrap','w:overflowPunct','w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN','w:bidi') + _AFTER_BIDI

def norm(t):
    return t.replace('ي', 'ی').replace('ك', 'ک')

def set_rtl_para(p, shade_fill=None):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1')
    pPr.insert_element_before(bidi, *_AFTER_BIDI)
    if shade_fill:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), shade_fill)
        pPr.insert_element_before(shd, *_AFTER_SHD)

def rtl_run(p, text, size=11, bold=False, italic=False, color=DARK):
    r = p.add_run(norm(text))
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(*color)
    rPr = r._r.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts'); rPr.insert_element_before(rF, 'w:b','w:bCs','w:i','w:iCs','w:color','w:sz','w:szCs','w:rtl')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rF.set(qn(a), FONT)
    if bold:
        bCs = OxmlElement('w:bCs'); rPr.insert_element_before(bCs, *_AFTER_BCS)
    if italic:
        iCs = OxmlElement('w:iCs'); rPr.insert_element_before(iCs, *_AFTER_ICS)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size * 2)))
    rPr.insert_element_before(szCs, *_AFTER_SZCS)
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1')
    rPr.insert_element_before(rtl, *_AFTER_RTL)
    return r

def para(doc, align=None, fill=None, r_ind=None, l_ind=None, before=0, after=4, line=1.15):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if r_ind is not None: pf.right_indent = Cm(r_ind)
    if l_ind is not None: pf.left_indent = Cm(l_ind)
    if align is not None: p.alignment = align
    set_rtl_para(p, fill)
    return p

def add_rich(p, text, size=11, color=DARK, bold_color=NAVY):
    """Parse **bold** and *italic* inline."""
    for part in re.split(r'(\*\*.+?\*\*|\*[^*]+\*)', text):
        if not part:
            continue
        bold = italic = False
        t = part
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            t = part[2:-2]; bold = True; color_k = bold_color
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            t = part[1:-1]; italic = True; color_k = color
        else:
            color_k = color
        rtl_run(p, t, size=size, bold=bold, italic=italic, color=color_k)

# ---------------------------------------------------------------- document
doc = Document()

# page setup A4
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.1)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)

# Normal style defaults
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(11)
st_rPr = st.element.get_or_add_rPr()
st_rF = st_rPr.find(qn('w:rFonts'))
if st_rF is None:
    st_rF = OxmlElement('w:rFonts'); st_rPr.insert(0, st_rF)
for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
    st_rF.set(qn(a), FONT)

# footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_rtl_para(fp)
rtl_run(fp, 'بسته تمرین‌های شخصی‌شده · عرفان کمال‌زاده', size=8.5, color=GRAY)

md = open(MD, encoding='utf-8').read()
lines = md.split('\n')

# ---------------------------------------------------------------- cover
def cover():
    para(doc, after=2)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    rtl_run(p, 'بسته تمرین‌های شخصی‌شده', size=21, bold=True, color=NAVY)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    rtl_run(p, 'عرفان کمال‌زاده · ۲۷ ساله', size=14, bold=True, color=DARK)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    rtl_run(p, '۱۰ تمرین فرویدی / ACT برای طرحواره‌های فعال', size=12.5, bold=True, color=BLUE)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    rtl_run(p, '(محرومیت هیجانی · معیارهای سرسختانه · تنبیه‌گری · بی‌اعتمادی · تسلیم · استحقاق)', size=10, color=GRAY)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_COVER, after=1, before=6)
    rtl_run(p, 'شناسه پرونده: RPT-INI-000014', size=10.5, color=DARK)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_COVER, after=10)
    rtl_run(p, 'تاریخ تنظیم: ۱۵ شهریور ۱۴۰۵', size=10.5, color=DARK)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_COVER, after=1)
    rtl_run(p, 'تهیه‌شده توسط: ', size=10.5, color=DARK)
    rtl_run(p, 'دکتر جواد میربلوکی', size=10.5, bold=True, color=NAVY)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_COVER, after=1)
    rtl_run(p, 'وب‌سایت و پلنر تمرین روزانه: ', size=10.5, color=DARK)
    rtl_run(p, 'joma.mirbolouki.com', size=10.5, bold=True, color=GOLDD)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_COVER, after=14)
    rtl_run(p, 'پشتیبانی (پیامک و پیام‌رسان بله): ', size=10.5, color=DARK)
    rtl_run(p, '۰۹۹۶۷۹۷۹۴۷۱', size=10.5, bold=True, color=NAVY)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    rtl_run(p, 'این بسته مکمل جلسهٔ روان‌درمانی است، نه جایگزین آن. هر هفته فقط یک یا دو تمرین انتخاب کنید.',
            size=9, italic=True, color=GRAY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

cover()

# ---------------------------------------------------------------- body state machine
fe_fix = {'حمام آخر شب دوش': 'حمام آخر شب 🚿'}
i = 0
started = False
while i < len(lines):
    s = lines[i].strip()
    if not s or s == '---':
        i += 1; continue
    if s.startswith('# بسته'):
        started = True; i += 1; continue

    if s.startswith('## تمرین'):
        m = re.match(r'^## تمرین (\S+)\s*—\s*(.+?)\s*\(([^)]+)\)\s*$', s)
        if m:
            title, tm = m.group(2).strip(), m.group(3).strip()
            for k, v in fe_fix.items():
                title = title.replace(k, v)
            p = para(doc, fill=FILL_TITLE, before=14, after=6)
            rtl_run(p, f'تمرین {m.group(1)} — {title}', size=13, bold=True, color=NAVY)
            rtl_run(p, f'  ({tm})', size=10.5, bold=False, color=GOLDD)
        else:
            p = para(doc, fill=FILL_TITLE, before=14, after=6)
            rtl_run(p, s[3:], size=13, bold=True, color=NAVY)
        i += 1; continue

    if s.startswith('## ') or s.startswith('# '):
        p = para(doc, fill=FILL_SECT, before=14, after=6)
        rtl_run(p, s.lstrip('# '), size=13, bold=True, color=GOLDD if False else NAVY)
        i += 1; continue

    if s.startswith('>'):
        grp = []
        while i < len(lines) and lines[i].strip().startswith('>'):
            grp.append(lines[i].strip()[1:].strip())
            i += 1
        if any('یادداشت برای شما' in g for g in grp):
            fill, sz, al = FILL_NOTE, 10.5, None
        elif any('هشدار بالینی' in g for g in grp):
            fill, sz, al = FILL_WARN, 10.5, None
        elif any('جوما ابزار' in g for g in grp):
            fill, sz, al = FILL_DISC, 9.5, WD_ALIGN_PARAGRAPH.CENTER
        else:  # exercise sub-quote / golden sentence
            fill, sz, al = (None, 11, None)
        first = True
        for g in grp:
            if not g:
                continue
            if '«آنچه من فکر می‌کنم' in g:
                p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, fill=FILL_SECT, before=4, after=1)
                rtl_run(p, g.replace('**', ''), size=12.5, bold=True, color=NAVY)
                continue
            p = para(doc, align=al, fill=fill, r_ind=0.6 if fill is None else None,
                     after=3, before=2 if first else 0)
            add_rich(p, g, size=sz)
            first = False
        continue

    if s.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        pf = p.paragraph_format
        pf.space_after = Pt(2); pf.line_spacing = 1.15; pf.right_indent = Cm(0.55)
        set_rtl_para(p)
        add_rich(p, s[2:], size=11)
        i += 1; continue

    mnum = re.match(r'^(\d+)\.\s+(.*)$', s)
    if mnum:
        p = para(doc, r_ind=0.55, after=2)
        rtl_run(p, f'{mnum.group(1)}. ', size=11, bold=True, color=BLUE)
        add_rich(p, mnum.group(2), size=11)
        i += 1; continue

    if s.startswith('*') and s.endswith('*') and not s.startswith('**'):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=0)
        rtl_run(p, s.strip('*'), size=9.5, italic=True, color=GRAY)
        i += 1; continue

    # labeled / plain paragraph
    p = para(doc, after=4)
    add_rich(p, s, size=11)
    i += 1

doc.save(OUT)
print('saved', OUT)

# quick sanity: re-open
d2 = Document(OUT)
n_p = len(d2.paragraphs)
bidi = d2.element.body.xpath('.//w:bidi')
rtl = d2.element.body.xpath('.//w:rtl')
print(f'paragraphs={n_p} bidi={len(bidi)} rtl_runs={len(rtl)}')
